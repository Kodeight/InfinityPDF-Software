import sys
import os
import json
import uuid
import tempfile
import random
import string
import re
from pathlib import Path
import traceback

# Library imports (win32com, reportlab, pypdf)
try:
    import win32com.client 
    import pythoncom
    from reportlab.pdfgen import canvas
    from reportlab.lib.colors import Color
    from pypdf import PdfReader, PdfWriter
    from pypdf.constants import UserAccessPermissions
except ImportError as e:
    print(json.dumps({"success": False, "error": f"Missing dependency: {e}"}))
    sys.exit(1)

def generate_owner_password():
    return ''.join(random.choices(string.ascii_letters + string.digits, k=16))

class InfinityBackend:
    """
    Core engine for InfinityPDF local version.
    Refactored to use win32com (PPTX) and pypdf (Security/Watermarking)
    to avoid heavy dependencies like pikepdf/LibreOffice.
    """
    
    def __init__(self, output_dir="InfinityPDF_Output"):
        self.output_dir = output_dir
        if not os.path.exists(self.output_dir):
            os.makedirs(self.output_dir)

    def apply_pdf_restrictions(self, input_path, permissions, output_path=None):
        """
        Applies PDF owner passwords and permission bitmasks using pypdf.
        permissions: dict with keys [print, copy, edit, restructure, forms, comments]
        output_path: Optional explicit output path. If not provided, generates one.
        """
        try:
            if output_path is None:
                output_path = os.path.join(self.output_dir, os.path.basename(input_path))
            
            reader = PdfReader(input_path)
            writer = PdfWriter()

            # Copy all pages
            for page in reader.pages:
                writer.add_page(page)

            # Map permissions
            # pypdf permissions are: PRINT, MODIFY, EXTRACT, etc.
            # Start with 0 (no permissions) if we want strict control, 
            # or just add strict flags.
            
            perms = UserAccessPermissions(0)
            
            # Note: The UI sends 'true' if the action is ALLOWED.
            if permissions.get('print', False):
                perms |= UserAccessPermissions.PRINT
            if permissions.get('copy', False):
                perms |= UserAccessPermissions.EXTRACT
            if permissions.get('edit', False):
                perms |= UserAccessPermissions.MODIFY
            if permissions.get('restructure', False):
                perms |= UserAccessPermissions.ASSEMBLE_DOC
            if permissions.get('forms', False):
                perms |= UserAccessPermissions.FILL_FORM_FIELDS
            if permissions.get('comments', False):
                perms |= UserAccessPermissions.ADD_OR_MODIFY

            owner_pw = generate_owner_password()

            # Encrypt with AES-128 for better compatibility with modern viewers
            # permissions_flag expects the integer value of flags
            writer.encrypt(
                user_password="", 
                owner_password=owner_pw, 
                permissions_flag=perms,
                algorithm="AES-128"
            )

            with open(output_path, "wb") as f:
                writer.write(f)

            return {"success": True, "path": output_path, "owner_token": owner_pw}

        except Exception as e:
            traceback.print_exc()
            return {"success": False, "error": str(e)}

    def batch_process_document(self, source_path, recipients, watermark_options, naming_pattern=None):
        """
        Converts PPTX to PDF (if needed) and applies watermarks per recipient.
        
        Creates a subfolder named after the source file.
        Supports naming_pattern with {{name}} placeholder for smart file naming.
        """
        results = []
        
        # Get base name from source file for subfolder
        base_name = Path(source_path).stem
        
        # Create subfolder for outputs
        output_subfolder = os.path.join(self.output_dir, base_name)
        os.makedirs(output_subfolder, exist_ok=True)
        
        # 1. Prepare base PDF
        base_pdf_path = os.path.join(tempfile.gettempdir(), f"base_{uuid.uuid4().hex}.pdf")
        
        try:
            if Path(source_path).suffix.lower() == '.pdf':
                import shutil
                shutil.copy2(source_path, base_pdf_path)
            else:
                self._pptx_to_pdf_win32(source_path, base_pdf_path)
        except Exception as e:
            traceback.print_exc()
            return [{"name": "Global Error", "status": "error", "error": f"Base PDF preparation failed: {e}"}]

        # Ensure pdfRestrictions dict exists (frontend now sends all flags explicitly)
        if 'pdfRestrictions' not in watermark_options:
            watermark_options['pdfRestrictions'] = {}

        custom_recipient = watermark_options.get('customWatermarkText', '').strip()
        recipients_to_process = [name for name in recipients if (name or "").strip()] if custom_recipient else list(recipients)
        if custom_recipient:
            recipients_to_process.append(custom_recipient)
            watermark_options['customWatermarkText'] = ''

        # Track used filenames to prevent duplicates
        used_filenames = set()
        total = len(recipients_to_process)

        # 2. Process for each recipient
        for i, name in enumerate(recipients_to_process):
            try:
                # Report progress per recipient
                progress = 5 + int(((i + 1) / max(total, 1)) * 90)
                print(f"PROGRESS:{progress}")
                
                # Handle empty/whitespace-only names gracefully
                clean_name = (name or "").strip()
                
                # If name is empty (single file mode), just use base name
                if not clean_name:
                    safe_name = base_name
                    if total > 1:
                        # Multiple recipients but one is empty — warn and use fallback
                        print(f"STATUS:Warning: Empty recipient at index {i+1}, using index-based name")
                        safe_name = f"{base_name}_{i+1}"
                else:
                    # Determine filename from naming pattern or default
                    if naming_pattern and naming_pattern.strip():
                        # Resolve {{name}} placeholder in the pattern
                        resolved = naming_pattern.replace('{{name}}', clean_name)
                        safe_name = self._sanitize_filename_preserve_punctuation(resolved)
                    else:
                        # Default naming: "basename - recipientname"
                        safe_name = self._sanitize_filename_preserve_punctuation(f"{base_name} - {clean_name}")
                
                # Handle duplicate filenames by appending index
                original_safe_name = safe_name
                counter = 2
                while safe_name in used_filenames:
                    safe_name = f"{original_safe_name}_{counter}"
                    counter += 1
                used_filenames.add(safe_name)
                
                output_filename = f"{safe_name}.pdf"
                output_path = os.path.join(output_subfolder, output_filename)
                
                print(f"STATUS:Processing {clean_name or f'file {i+1}'}...")
                
                # Apply watermark and save with restrictions
                self._add_watermark_and_save(base_pdf_path, output_path, clean_name, watermark_options)
                
                results.append({"name": clean_name or f"file_{i+1}", "status": "success", "path": output_path})
            except Exception as e:
                traceback.print_exc()
                results.append({"name": name or f"file_{i+1}", "status": "error", "error": str(e)})
        
        # Cleanup base PDF
        if os.path.exists(base_pdf_path):
            os.remove(base_pdf_path)
        
        # Summary
        success_count = sum(1 for r in results if r.get("status") == "success")
        error_count = sum(1 for r in results if r.get("status") == "error")
        if error_count > 0:
            print(f"STATUS:Completed with {error_count} error(s) out of {total}")
            
        return results

    def _sanitize_filename_preserve_punctuation(self, name):
        """
        Sanitize filename but preserve common punctuation like (), -, _.
        Only removes characters that are invalid on Windows filesystem.
        """
        # Invalid Windows filename chars: \ / : * ? " < > |
        invalid_chars = r'\/:*?"<>|'
        result = ''.join(c for c in name if c not in invalid_chars)
        return result.strip()

    def _pptx_to_pdf_win32(self, pptx_path, output_pdf_path):
        pythoncom.CoInitialize()
        powerpoint = None
        prs = None
        try:
            pptx_abs_path = os.path.abspath(pptx_path)
            pdf_abs_path = os.path.abspath(output_pdf_path)

            powerpoint = win32com.client.Dispatch("PowerPoint.Application")
            prs = powerpoint.Presentations.Open(pptx_abs_path, WithWindow=False)
            
            # 32 = ppSaveAsPDF
            prs.SaveAs(pdf_abs_path, 32)
        finally:
            if prs: prs.Close()
            # Don't quit PowerPoint if user has it open? 
            # Ideally quit only if we started it, but simple approach:
            if powerpoint: powerpoint.Quit()
            pythoncom.CoUninitialize()

    def _add_watermark_and_save(self, input_pdf, output_pdf, watermark_text, options):
        """
        Adds watermark to PDF and applies permissions.
        Integrated from convert.py logic for full feature parity.
        If options.enabled is False, watermark is skipped but restrictions still apply.
        """
        import io
        
        # Check if watermark is enabled (default True for backwards compat)
        watermark_enabled = options.get('enabled', True)
        
        # Parse watermark options with defaults (matching convert.py exactly)
        opacity = options.get('opacity', 20) / 100.0  
        font_size = options.get('fontSize', 36)
        rotation = options.get('rotation', -30)
        diagonal = options.get('diagonal', False)
        position_x_pct = options.get('positionX', 50) / 100.0
        position_y_pct = options.get('positionY', 50) / 100.0
        
        # PDF restrictions from options
        pdf_restrictions = options.get('pdfRestrictions', {})
        
        reader = PdfReader(input_pdf)
        writer = PdfWriter()

        for page in reader.pages:
            if watermark_enabled:
                # Create watermark layer
                packet = io.BytesIO()
                page_box = page.mediabox
                width = float(page_box.width)
                height = float(page_box.height)
                
                can = canvas.Canvas(packet, pagesize=(width, height))
                can.saveState()
                
                # Calculate position (matching convert.py logic exactly)
                if diagonal:
                    # If diagonal, force center it and use 45 degrees
                    x_pos = width / 2
                    y_pos = height / 2
                    rotation_to_use = 45
                else:
                    # Use custom position and rotation
                    x_pos = width * position_x_pct
                    y_pos = height * position_y_pct
                    rotation_to_use = rotation
                
                # Apply watermark color with opacity
                watermark_color = Color(0, 0, 0, alpha=opacity)
                
                can.translate(x_pos, y_pos)
                can.rotate(rotation_to_use)
                can.setFont("Helvetica-Bold", font_size)
                can.setFillColor(watermark_color)
                
                # Draw text centered (matching convert.py)
                text_width = can.stringWidth(watermark_text, "Helvetica-Bold", font_size)
                can.drawString(-text_width/2, -font_size/2, watermark_text)
                
                can.restoreState()
                can.save()
                packet.seek(0)
                
                watermark_pdf = PdfReader(packet)
                page.merge_page(watermark_pdf.pages[0])
            
            writer.add_page(page)

        # Generate owner password
        owner_pw = generate_owner_password()
        
        # Build permissions from pdfRestrictions (matching convert.py logic)
        perms = UserAccessPermissions(0)
        
        # Note: UI sends True if action is ALLOWED
        if pdf_restrictions.get('allowPrinting', False):
            perms |= UserAccessPermissions.PRINT
        if pdf_restrictions.get('allowCopying', False):
            perms |= UserAccessPermissions.EXTRACT
        if pdf_restrictions.get('allowEditing', False):
            perms |= UserAccessPermissions.MODIFY
        if pdf_restrictions.get('allowAssemble', False):
            perms |= UserAccessPermissions.ASSEMBLE_DOC
        if pdf_restrictions.get('allowFillForm', False):
            perms |= UserAccessPermissions.FILL_FORM_FIELDS
        if pdf_restrictions.get('allowAnnotate', False):
            perms |= UserAccessPermissions.ADD_OR_MODIFY

        # Encrypt with AES-128 and apply permissions
        try:
            writer.encrypt(
                user_password="", 
                owner_password=owner_pw, 
                permissions_flag=perms,
                algorithm="AES-128"
            )
        except TypeError:
            # Fallback for older pypdf versions
            try:
                writer.encrypt(user_password="", owner_password=owner_pw, permissions=perms)
            except Exception:
                writer.encrypt(user_password="", owner_password=owner_pw)
        
        with open(output_pdf, "wb") as f:
            writer.write(f)

    def universal_convert(self, input_path, target_format, output_dir=None):
        """
        Universal converter engine.
        Supports:
        - Office (DOCX, PPTX, XLSX) -> PDF
        - PDF -> DOCX
        - Images -> PDF
        """
        try:
            if output_dir is None:
                output_dir = os.path.join(self.output_dir, "Universal")
            
            if not os.path.exists(output_dir):
                os.makedirs(output_dir)

            input_path = os.path.abspath(input_path)
            base_name = os.path.splitext(os.path.basename(input_path))[0]
            ext = os.path.splitext(input_path)[1].lower()
            target_format = target_format.upper()
            
            output_file = f"{base_name}.{target_format.lower()}"
            output_path = os.path.join(output_dir, output_file)
            
            # Strategy Selection
            if target_format == "PDF":
                if ext in ['.pptx', '.ppt']:
                    self._pptx_to_pdf_win32(input_path, output_path)
                elif ext in ['.docx', '.doc']:
                    from docx2pdf import convert
                    convert(input_path, output_path)
                elif ext in ['.xlsx', '.xls']:
                    self._excel_to_pdf_win32(input_path, output_path)
                elif ext in ['.jpg', '.jpeg', '.png', '.bmp', '.tiff']:
                    from PIL import Image
                    image = Image.open(input_path)
                    if image.mode != 'RGB':
                        image = image.convert('RGB')
                    image.save(output_path)
                else:
                    return {"success": False, "error": f"Unsupported input format for PDF conversion: {ext}"}
            
            elif target_format in ["DOCX", "DOC"]:
                if ext == '.pdf':
                    try:
                        from pdf2docx import Converter
                        # pdf2docx only supports .docx output, not legacy .doc format
                        # Force output to .docx regardless of requested format
                        docx_path = f"{os.path.splitext(output_path)[0]}.docx" if output_path.endswith('.doc') else output_path
                        if not docx_path.endswith('.docx'):
                            docx_path = docx_path.replace(os.path.splitext(docx_path)[1], '.docx')
                        
                        cv = Converter(input_path)
                        cv.convert(docx_path)
                        cv.close()
                        if self._pdf_contains_arabic(input_path):
                            self._normalize_docx_arabic_rtl(docx_path)
                        
                        # Update output_path to the actual created file
                        output_path = docx_path
                    except ImportError:
                        return {"success": False, "error": "pdf2docx library not installed. Install with: pip install pdf2docx"}
                    except Exception as e:
                        if self._pdf_contains_arabic(input_path):
                            try:
                                self._pdf_to_docx_arabic_text(input_path, docx_path)
                                output_path = docx_path
                            except Exception:
                                return {"success": False, "error": f"PDF to DOCX conversion failed: {str(e)}"}
                        else:
                            return {"success": False, "error": f"PDF to DOCX conversion failed: {str(e)}"}
                else:
                    return {"success": False, "error": f"Unsupported input format for Word conversion: {ext}"}
            
            elif target_format in ["PNG", "JPG"]:
                if ext == '.pdf':
                    output_path = self._pdf_to_image(input_path, output_path, target_format)
                elif ext in ['.jpg', '.jpeg', '.png', '.bmp', '.tiff']:
                    from PIL import Image
                    image = Image.open(input_path)
                    if target_format == "JPG" and image.mode in ("RGBA", "P"):
                        image = image.convert("RGB")
                    image.save(output_path)
                elif ext in ['.docx', '.doc', '.pptx', '.ppt', '.xlsx', '.xls']:
                    # Double conversion: Office -> PDF -> Image
                    import shutil
                    temp_subfolder = os.path.join(tempfile.gettempdir(), f"ipdf_temp_{uuid.uuid4().hex}")
                    os.makedirs(temp_subfolder, exist_ok=True)
                    
                    try:
                        conv_result = self.universal_convert(input_path, "PDF", temp_subfolder)
                        if not conv_result.get("success"):
                            return conv_result
                        
                        actual_temp_pdf = conv_result["path"]
                        output_path = self._pdf_to_image(actual_temp_pdf, output_path, target_format)
                    finally:
                        shutil.rmtree(temp_subfolder, ignore_errors=True)
                else:
                    return {"success": False, "error": f"Unsupported input format for Image conversion: {ext}"}

            elif target_format == "PPTX":
                if ext == '.pdf':
                    self._pdf_to_pptx(input_path, output_path)
                elif ext in ['.jpg', '.jpeg', '.png', '.bmp', '.tiff']:
                    self._image_to_pptx([input_path], output_path)
                elif ext in ['.docx', '.doc', '.xlsx', '.xls']:
                    # Double conversion: Office -> PDF -> PPTX
                    import shutil
                    temp_subfolder = os.path.join(tempfile.gettempdir(), f"ipdf_temp_{uuid.uuid4().hex}")
                    os.makedirs(temp_subfolder, exist_ok=True)
                    
                    try:
                        conv_result = self.universal_convert(input_path, "PDF", temp_subfolder)
                        if not conv_result.get("success"):
                            return conv_result
                        
                        actual_temp_pdf = conv_result["path"]
                        self._pdf_to_pptx(actual_temp_pdf, output_path)
                    finally:
                        shutil.rmtree(temp_subfolder, ignore_errors=True)
                else:
                    return {"success": False, "error": f"Unsupported input format for PPTX conversion: {ext}"}

            else:
                 return {"success": False, "error": f"Unsupported target format: {target_format}"}

            return {"success": True, "path": output_path, "outputDir": output_dir}

        except Exception as e:
            traceback.print_exc()
            return {"success": False, "error": str(e)}

    def _pdf_to_image(self, pdf_path, output_path, format):
        import fitz  # PyMuPDF
        doc = fitz.open(pdf_path)
        # We only take the first page for "Universal Converter" simplicity 
        # unless user wants a zip of all pages, but usually "Convert to Image"
        # for a multi-page doc is better handled as single output.
        # If the user wants ALL pages, we should create a folder.
        # For now, let's just do page 1 or append page number if multi-page.
        
        if len(doc) > 1:
            base, ext = os.path.splitext(output_path)
            for i, page in enumerate(doc):
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2)) # 2x scale for quality
                p_output = f"{base}_page_{i+1}{ext}"
                pix.save(p_output)
            # Final path for UI display can be the first one
            output_path = f"{base}_page_1{ext}"
        else:
            page = doc[0]
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
            pix.save(output_path)
        doc.close()
        return output_path

    def _contains_arabic(self, text):
        return bool(re.search(r'[\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF]', text or ""))

    def _pdf_contains_arabic(self, pdf_path):
        import fitz
        doc = fitz.open(pdf_path)
        try:
            return any(self._contains_arabic(page.get_text("text")) for page in doc)
        finally:
            doc.close()

    def _set_docx_rtl(self, paragraph):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        p_pr = paragraph._p.get_or_add_pPr()
        bidi = p_pr.find(qn('w:bidi'))
        if bidi is None:
            bidi = OxmlElement('w:bidi')
            p_pr.append(bidi)
        bidi.set(qn('w:val'), '1')

        for run in paragraph.runs:
            r_pr = run._r.get_or_add_rPr()
            rtl = r_pr.find(qn('w:rtl'))
            if rtl is None:
                rtl = OxmlElement('w:rtl')
                r_pr.append(rtl)
            rtl.set(qn('w:val'), '1')

    def _normalize_docx_arabic_rtl(self, docx_path):
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document(docx_path)
        changed = False

        for paragraph in doc.paragraphs:
            if self._contains_arabic(paragraph.text):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                self._set_docx_rtl(paragraph)
                changed = True

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        if self._contains_arabic(paragraph.text):
                            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                            self._set_docx_rtl(paragraph)
                            changed = True

        if changed:
            doc.save(docx_path)

    def _pdf_to_docx_arabic_text(self, pdf_path, output_path):
        import fitz
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        pdf = fitz.open(pdf_path)
        doc = Document()
        style = doc.styles['Normal']
        style.font.name = 'Arial'

        try:
            for page_index, page in enumerate(pdf):
                if page_index:
                    doc.add_page_break()

                blocks = page.get_text("blocks")
                for block in sorted(blocks, key=lambda b: (round(b[1], 1), b[0])):
                    text = block[4].strip()
                    if not text:
                        continue

                    paragraph = doc.add_paragraph()
                    is_arabic = self._contains_arabic(text)
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if is_arabic else WD_ALIGN_PARAGRAPH.LEFT
                    run = paragraph.add_run(text)
                    run.font.name = 'Arial'
                    if is_arabic:
                        self._set_docx_rtl(paragraph)

            doc.save(output_path)
        finally:
            pdf.close()

    def _pdf_to_pptx(self, pdf_path, output_path):
        import fitz
        doc = fitz.open(pdf_path)
        temp_dir = os.path.join(tempfile.gettempdir(), f"ipdf_pdf2pptx_{uuid.uuid4().hex}")
        os.makedirs(temp_dir, exist_ok=True)
        
        image_paths = []
        try:
            for i, page in enumerate(doc):
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
                img_path = os.path.join(temp_dir, f"page_{i}.png")
                pix.save(img_path)
                image_paths.append(img_path)
                
            self._image_to_pptx(image_paths, output_path)
        finally:
            import shutil
            shutil.rmtree(temp_dir, ignore_errors=True)
            doc.close()

    def _image_to_pptx(self, image_paths, output_path):
        from pptx import Presentation
        
        prs = Presentation()
        blank_slide_layout = prs.slide_layouts[6]
        
        for img_path in image_paths:
            slide = prs.slides.add_slide(blank_slide_layout)
            try:
                # Fill slide
                slide.shapes.add_picture(img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)
            except Exception as e:
                print(f"Failed to add image to slide: {e}")
                
        prs.save(output_path)

    def _excel_to_pdf_win32(self, input_path, output_path):
        pythoncom.CoInitialize()
        excel = None
        wb = None
        try:
            excel = win32com.client.Dispatch("Excel.Application")
            excel.Interactive = False
            excel.Visible = False
            wb = excel.Workbooks.Open(input_path)
            # 0 = xlTypePDF
            wb.ExportAsFixedFormat(0, output_path)
        finally:
            if wb: wb.Close(False)
            if excel: excel.Quit()
            pythoncom.CoUninitialize()

if __name__ == "__main__":
    if len(sys.argv) > 1:
        command = sys.argv[1]
        
        if command == 'apply-security' and len(sys.argv) >= 5:
            input_path = sys.argv[2]
            permissions_json = sys.argv[3]
            output_path = sys.argv[4]
            
            try:
                permissions = json.loads(permissions_json)
                engine = InfinityBackend(os.path.dirname(output_path) or ".")
                result = engine.apply_pdf_restrictions(input_path, permissions, output_path)
                print(json.dumps(result))
                sys.exit(0 if result.get('success') else 1)
            except Exception as e:
                traceback.print_exc()
                print(json.dumps({"success": False, "error": str(e)}))
                sys.exit(1)
        
        elif command == 'multi-pdf' and len(sys.argv) >= 6:
            source_path = sys.argv[2]
            output_dir = sys.argv[3]
            student_list_json = sys.argv[4]
            watermark_options_json = sys.argv[5]
            naming_pattern = sys.argv[6] if len(sys.argv) > 6 else None
            
            try:
                student_list = json.loads(student_list_json)
                watermark_options = json.loads(watermark_options_json)
                
                # Validate student list structure
                students = student_list.get('students', [])
                if not isinstance(students, list):
                    print(json.dumps({"success": False, "error": "Invalid student list format", "errorType": "INVALID_DATA"}))
                    sys.exit(1)
                
                # Extract recipient names from student list
                recipients = [s.get('name', '') for s in students]
                has_custom_recipient = bool(watermark_options.get('customWatermarkText', '').strip())
                
                if len(recipients) == 0 and not has_custom_recipient:
                    print(json.dumps({"success": False, "error": "No recipients provided", "errorType": "EMPTY_LIST"}))
                    sys.exit(1)
                
                total_recipients = len([name for name in recipients if (name or "").strip()]) + (1 if has_custom_recipient else 0)
                print(f"STATUS:Starting batch conversion for {total_recipients or len(recipients)} recipients...")
                print("PROGRESS:5")
                
                engine = InfinityBackend(output_dir)
                results = engine.batch_process_document(source_path, recipients, watermark_options, naming_pattern)
                
                print("PROGRESS:100")
                print(json.dumps({"success": True, "results": results}))
                sys.exit(0)
            except json.JSONDecodeError as e:
                print(json.dumps({"success": False, "error": f"Invalid JSON input: {e}", "errorType": "INVALID_JSON"}))
                sys.exit(1)
            except Exception as e:
                traceback.print_exc()
                print(json.dumps({"success": False, "error": str(e)}))
                sys.exit(1)

        elif command == 'universal-convert' and len(sys.argv) >= 4:
            input_path = sys.argv[2]
            target_format = sys.argv[3]
            output_dir = sys.argv[4] if len(sys.argv) > 4 else None
            
            try:
                 engine = InfinityBackend(output_dir if output_dir else "InfinityPDF_Output")
                 result = engine.universal_convert(input_path, target_format, output_dir)
                 print(json.dumps(result))
                 sys.exit(0 if result.get('success') else 1)
            except Exception as e:
                traceback.print_exc()
                print(json.dumps({"success": False, "error": str(e)}))
                sys.exit(1)
        
        else:
            print(json.dumps({"success": False, "error": "Invalid command"}))
            sys.exit(1)
    else:
        print("InfinityPDF Backend Engine (pypdf/win32)")
