#!/usr/bin/env python3
"""PDF TO WORD - convert PDF to DOCX with RTL normalization (new tool, additive).

CLI: python pdf_to_word.py <operation> <args-json> <output-dir>
Ops:
  convert(pdf) -> .docx via pdf2docx, then own RTL normalization
  with python-docx (right-align + w:bidi/w:rtl on Arabic paragraphs).

If pdf2docx raises, success False with its message is returned (no crash).
Stdout stays PROGRESS:/RESULT: only, so converter chatter is silenced.
"""
import os, sys; sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import contextlib
import io
import re

from _common import (
    check_cancel, log, progress, require_arg,
    run_tool, unique_path, validate_pdf,
)

ARABIC_RE = re.compile(r"[\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF\uFB50-\uFDFF\uFE70-\uFEFF]")


def _rtl_paragraph(paragraph):
    """Right-align + bidi/rtl markup for one paragraph. Returns True if applied."""
    try:
        text = paragraph.text or ""
    except Exception:
        return False
    if not ARABIC_RE.search(text):
        return False
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    try:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    except Exception:
        pass
    try:
        pPr = paragraph._p.get_or_add_pPr()
        if pPr.find(qn("w:bidi")) is None:
            bidi = OxmlElement("w:bidi")
            pPr.append(bidi)
    except Exception as e:
        log("bidi markup failed: %s" % e)
    for run in paragraph.runs:
        try:
            rPr = run._r.get_or_add_rPr()
            if rPr.find(qn("w:rtl")) is None:
                rtl = OxmlElement("w:rtl")
                rPr.append(rtl)
        except Exception:
            continue
    return True


def _rtl_normalize(docx_path):
    from docx import Document
    doc = Document(docx_path)
    count = 0
    for p in doc.paragraphs:
        check_cancel()
        if _rtl_paragraph(p):
            count += 1
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if _rtl_paragraph(p):
                        count += 1
    doc.save(docx_path)
    return count


def op_convert(args, outdir):
    pdf = validate_pdf(require_arg(args, "pdf"))
    try:
        from pdf2docx import Converter
    except Exception as e:
        return {"success": False, "error": "pdf2docx import failed: %s" % e,
                "outputs": [], "info": {}}
    base = os.path.splitext(os.path.basename(pdf))[0].strip() or "document"
    out = unique_path(outdir, base, "docx")
    progress(5)
    try:
        cv = Converter(pdf)
        try:
            with contextlib.redirect_stdout(io.StringIO()):
                cv.convert(out)
        finally:
            try:
                cv.close()
            except Exception:
                pass
    except Exception as e:
        log("pdf2docx conversion failed: %s" % e)
        try:
            if os.path.exists(out):
                os.remove(out)
        except Exception:
            pass
        return {"success": False, "error": "pdf2docx conversion failed: %s" % e,
                "outputs": [], "info": {"file": os.path.basename(pdf)}}
    progress(85)
    check_cancel()
    try:
        rtl_count = _rtl_normalize(out)
    except Exception as e:
        log("RTL normalization failed: %s" % e)
        rtl_count = 0
    progress(100)
    return {"outputs": [out], "info": {
        "file": os.path.basename(pdf),
        "rtl_paragraphs": rtl_count,
        "output_bytes": os.path.getsize(out) if os.path.exists(out) else 0,
    }}


OPS = {"convert": op_convert}

if __name__ == "__main__":
    run_tool(OPS)
