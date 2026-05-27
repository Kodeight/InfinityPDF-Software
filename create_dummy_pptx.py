
from pptx import Presentation
import os

def create_dummy_pptx(filename="dummy.pptx"):
    prs = Presentation()
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]

    title.text = "Hello, World!"
    subtitle.text = "This is a test PPTX for InfinityPDF."

    prs.save(filename)
    print(f"Created {filename}")

if __name__ == "__main__":
    create_dummy_pptx()
